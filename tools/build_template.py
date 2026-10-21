#!/usr/bin/env python3
"""Costruisce il template docxtemplater del modulo di autorizzazione al pagamento
a partire dal modello aziendale DOCX (tools/assets).
Uso: python3 tools/build_template.py [sorgente.docx] [destinazione.docx]
"""
import copy, sys, os
import docx
from docx.shared import Pt

SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), 'assets', 'Autorizzazione_congiunta_pagamenti_CEO_CFO_modello_v1.docx')
DST = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(__file__), '..', 'server', 'src', 'templates', 'autorizzazione_pagamento.docx')

d = docx.Document(SRC)

def set_text(p, text):
    """Sostituisce il testo del paragrafo con un solo run (mantiene la formattazione del primo run)."""
    runs = p.runs
    first = runs[0]._element if runs else None
    for el in list(p._p):
        if el.tag.endswith('}pPr') or el is first:
            continue
        p._p.remove(el)   # altri run, revisioni (w:ins), smartTag, hyperlink
    if runs:
        runs[0].text = text
    else:
        p.add_run(text)

def cell(t, r, c, text, pi=0):
    set_text(d.tables[t].rows[r].cells[c].paragraphs[pi], text)

paras = d.paragraphs
# Titolo e riferimento normativo (modificabili dalle Impostazioni)
set_text(paras[0], '{title}')
set_text(paras[2], '{form_ref}')

# Tabella 0: protocollo
cell(0, 0, 1, '{protocol}'); cell(0, 0, 3, '{request_date}')
cell(0, 1, 1, '{requester}'); cell(0, 1, 3, '{cost_center}')
# Tabella 1: dati del pagamento
cell(1, 0, 1, '{beneficiary_name}'); cell(1, 1, 1, '{beneficiary_tax_id}')
cell(1, 2, 1, '{beneficiary_iban}'); cell(1, 3, 1, '{beneficiary_bank}')
cell(1, 4, 1, '€ {amount}'); cell(1, 5, 1, 'Euro {amount_words}')
cell(1, 6, 1, '{purpose}')
cell(1, 7, 1, '{chk_fattura} Fattura   {chk_contratto} Contratto   {chk_ordine} Ordine   {chk_doc_altro} Altro: {doc_other}   n. {doc_number} del {doc_date}')
cell(1, 8, 1, '{debit_account}')
cell(1, 9, 1, '{chk_sepa} Bonifico SEPA   {chk_urgente} Bonifico urgente   {chk_estero} Bonifico estero   {chk_f24} F24   {chk_pm_altro} Altro: {pm_other}')
cell(1, 10, 1, '{value_date}')
# Tabella 2: scaglione. Aggiungo in testa la riga "fino a 50.000" (firma singola)
t2 = d.tables[2]
row_joint = t2.rows[1]
new_tr = copy.deepcopy(row_joint._tr)
row_joint._tr.addprevious(new_tr)
set_text(t2.rows[1].cells[0].paragraphs[0], '{chk_tier_single}')
set_text(t2.rows[1].cells[1].paragraphs[0], 'Fino a {threshold_single}')
set_text(t2.rows[1].cells[2].paragraphs[0], 'Operazione rientrante nel Piano Industriale, nel Piano Finanziario e nel budget approvati, o derivante da obblighi già autorizzati')
set_text(t2.rows[1].cells[3].paragraphs[0], 'Firma singola dell’AD o del CFO: {single_signer_label}')
set_text(t2.rows[2].cells[0].paragraphs[0], '{chk_tier_joint}')
set_text(t2.rows[2].cells[1].paragraphs[0], 'Oltre {threshold_single} e fino a {threshold_joint}')
set_text(t2.rows[3].cells[0].paragraphs[0], '{chk_tier_cda}')
set_text(t2.rows[3].cells[1].paragraphs[0], 'Oltre {threshold_joint}')
set_text(t2.rows[3].cells[2].paragraphs[0], 'Preventiva approvazione del Consiglio di Amministrazione – delibera n. {board_number} del {board_date}')
# Paragrafo "pagamento già approvato"
for p in d.paragraphs:
    if p.text.startswith('☐  Pagamento di paghe'):
        set_text(p, '{chk_prior}  Pagamento di paghe, imposte, contributi, rate di debito o altro pagamento già specificamente approvato (estremi approvazione: {prior_ref}): non è richiesta una nuova approvazione del CdA per la mera esecuzione; si applicano i poteri di firma corrispondenti all’importo.')
    elif p.text.startswith('Nota: le operazioni di importo unitario fino a'):
        set_text(p, 'Nota: le operazioni di importo unitario fino a {threshold_single} rientranti nel Piano Industriale, nel Piano Finanziario e nel budget approvati, o derivanti da obblighi già autorizzati, sono eseguibili con firma singola dell’AD o del CFO. È vietato il frazionamento artificioso delle operazioni al fine di eludere le soglie.')
    elif p.text.startswith('I sottoscritti, nelle rispettive qualità'):
        set_text(p, '{auth_text}')
    elif p.text.startswith('2. Scaglione applicabile'):
        set_text(p, '2. Scaglione applicabile')
    elif p.text.startswith('Il presente modulo, completo della documentazione'):
        last = p
# Tabella 3: verifiche preliminari
for i in range(6):
    cell(3, i, 0, '{chk_v%d}' % (i + 1))
# Tabella 4: firme
t4 = d.tables[4]
c_ceo, c_cfo = t4.rows[0].cells[0], t4.rows[0].cells[1]
set_text(c_ceo.paragraphs[1], '{ceo_name}'); set_text(c_ceo.paragraphs[2], 'Data {ceo_date}'); set_text(c_ceo.paragraphs[3], 'Firma: {ceo_signature}')
set_text(c_cfo.paragraphs[1], '{cfo_name}'); set_text(c_cfo.paragraphs[2], 'Data {cfo_date}'); set_text(c_cfo.paragraphs[3], 'Firma: {cfo_signature}')
# Tabella 5: esecuzione
cell(5, 0, 1, '{executed_by}'); cell(5, 0, 3, '{executed_at}')
cell(5, 1, 1, '{cro_trn}'); cell(5, 1, 3, '{attachments_count}')
# Riga finale di tracciabilità
foot = d.add_paragraph()
r = foot.add_run('{trace_line}')
r.font.size = Pt(7.5); r.italic = True; r.font.name = 'Arial'

os.makedirs(os.path.dirname(DST), exist_ok=True)
d.save(DST)
print('template scritto in', DST)
